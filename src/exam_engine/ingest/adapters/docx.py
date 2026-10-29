"""DOCX adapter (python-docx via the optional loader): body in file order, one logical page.

Reads only the body: paragraphs, headings, list items, tables, inline images and the plain text of
OMML math. Headers, footers, text boxes and footnotes are not read (warning `unsupported_content`).
The original file is only read; no network.
"""
import hashlib
import io
import os
import re
import unicodedata
import zipfile

from exam_engine.ingest import optional, segment
from exam_engine.ingest.schema import SCHEMA_VERSION, IngestError

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
M = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"  # password-protected docx is an OLE container
_HEADING = re.compile(r"^Heading\s+(\d+)$", re.I)
_MIME_EXT = {"image/png": ".png", "image/jpeg": ".jpg", "image/gif": ".gif",
             "image/bmp": ".bmp", "image/tiff": ".tiff"}


def _warn(code, msg, page=None):
    w = {"level": "warning", "code": code, "message": msg}
    if page is not None:
        w["page"] = page
    return w


def _nfc(s):
    return unicodedata.normalize("NFC", s)


def _text(para, state):
    """Paragraph text in document order incl. OMML text; text-box content is skipped."""
    out = []
    for el in para.iter():
        tag = el.tag
        if tag == W + "txbxContent":
            state["unsupported"].add("text box")
        if tag in (W + "t", M + "t"):
            if any(a.tag == W + "txbxContent" for a in el.iterancestors()):
                continue
            out.append(el.text or "")
            if tag == M + "t":
                state["math"] = True
        elif tag == W + "tab":
            out.append("\t")
        elif tag in (W + "br", W + "cr"):
            out.append("\n")
    return _nfc("".join(out)).strip()


def _style(para):
    try:
        return para.style.name or ""
    except Exception:  # noqa: BLE001 - a broken style reference is a plain paragraph
        return ""


def _kind(p_el, style):
    """(block type, heading level)."""
    m = _HEADING.match(style)
    if m:
        return "heading", int(m.group(1))
    if style.lower() == "title":
        return "heading", 1
    if style.lower().startswith("list") or p_el.find(".//" + W + "numPr") is not None:
        return "list", None
    return "paragraph", None


def _table_rows(tbl, state):
    rows, seen = [], []
    for row in tbl.rows:
        cells = []
        for c in row.cells:
            if any(c._tc is s for s in seen):  # merged cell seen once
                continue
            seen.append(c._tc)
            parts = [_text(p._p, state) for p in c.paragraphs]
            cells.append(" ".join(t for t in parts if t))
        rows.append(cells)
    return rows


def _unsupported_parts(doc, raw):
    found = []
    try:
        for sec in doc.sections:
            for label, part in (("header", sec.header), ("footer", sec.footer)):
                if any(p.text.strip() for p in part.paragraphs) and label not in found:
                    found.append(label)
    except Exception:  # noqa: BLE001 - best effort
        pass
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            if "word/footnotes.xml" in z.namelist() and re.search(
                    rb"<w:t[ >]", z.read("word/footnotes.xml")):
                found.append("footnote")
    except Exception:  # noqa: BLE001
        pass
    return found


def ingest(path, config=None):
    config = config or {}
    try:
        with open(path, "rb") as fh:
            raw = fh.read()
    except OSError as exc:
        raise IngestError("corrupt", "cannot read %s: %s" % (path, exc))
    name = os.path.basename(path)
    if not raw:
        raise IngestError("empty", "%s has no content" % name)
    if raw.startswith(OLE_MAGIC):
        raise IngestError("encrypted", "%s is password protected (or an old .doc)" % name)
    if not zipfile.is_zipfile(io.BytesIO(raw)):
        raise IngestError("corrupt", "%s is not a DOCX (zip) file" % name)
    docx = optional.load("docx", "python-docx")
    try:
        doc = docx.Document(io.BytesIO(raw))
    except Exception as exc:  # noqa: BLE001 - any parse failure is data
        raise IngestError("corrupt", "cannot open DOCX %s: %s" % (name, exc))
    return _extract(docx, doc, raw, name, config)


def _extract(docx, doc, raw, name, config):
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    sha = hashlib.sha256(raw).hexdigest()
    asset_dir = config.get("asset_dir")
    rel_dir = None
    if asset_dir:
        rel = os.path.relpath(asset_dir).replace(os.sep, "/")
        rel_dir = os.path.basename(os.path.normpath(asset_dir)) if rel.startswith("..") else rel
    state = {"unsupported": set(), "math": False}
    blocks, assets, warnings, lines = [], [], [], []

    def add(btype, text, level=None, asset=None):
        blocks.append({"id": "b%d" % (len(blocks) + 1), "order": len(blocks), "type": btype,
                       "text": text, "level": level, "bbox": None, "origin": "source",
                       "asset": asset})

    def add_image(blip):
        rid = blip.get(R + "embed")
        try:
            part = doc.part.related_parts[rid]
            blob = part.blob
        except Exception:  # noqa: BLE001 - external or dangling link
            warnings.append(_warn("asset_not_saved", "image relation %s not readable" % rid, 1))
            add("image", "")
            return
        if not asset_dir:
            warnings.append(_warn("asset_not_saved", "no --assets folder; image not saved", 1))
            add("image", "")
            return
        ext = os.path.splitext(str(part.partname))[1] or _MIME_EXT.get(part.content_type, ".bin")
        fname = "docx-img-%03d%s" % (len(assets) + 1, ext.lower())
        os.makedirs(asset_dir, exist_ok=True)
        with open(os.path.join(asset_dir, fname), "wb") as fh:
            fh.write(blob)
        aid = "a%d" % (len(assets) + 1)
        assets.append({"id": aid, "path": "/".join([rel_dir, fname]),
                       "mime": part.content_type, "sha256": hashlib.sha256(blob).hexdigest(),
                       "page": 1, "bbox": None})
        add("image", "", asset=aid)

    for el in doc.element.body.iterchildren():
        if el.tag == W + "p":
            para = Paragraph(el, doc)
            text = _text(el, state)
            if text:
                btype, level = _kind(el, _style(para))
                add(btype, text, level)
                lines.extend(text.split("\n"))
            for blip in el.iter(A + "blip"):
                add_image(blip)
        elif el.tag == W + "tbl":
            tbl = Table(el, doc)
            rows = _table_rows(tbl, state)
            if any(any(c for c in r) for r in rows):
                add("table", "\n".join(" | ".join(r) for r in rows))
                blocks[-1]["rows"] = rows
            if next(el.iter(A + "blip"), None) is not None:
                state["unsupported"].add("image in table")
    for label in _unsupported_parts(doc, raw):
        state["unsupported"].add(label)
    if not blocks:
        raise IngestError("empty", "%s has no text and no images" % name)
    pwarn = []
    if state["math"]:
        pwarn.append(_warn("math_possible", "equation text kept as plain text, not LaTeX", 1))
    if state["unsupported"]:
        pwarn.append(_warn("unsupported_content", "not read: %s"
                           % ", ".join(sorted(state["unsupported"])), 1))
    warnings = pwarn + warnings
    intro, qs, ok = segment.segment(lines)
    if not qs:
        warnings.append(_warn("no_question_numbering", "no question numbering found"))
    elif not ok:
        warnings.append(_warn("segmentation_skipped",
                              "question numbering not consecutive or a question has no text"))
    out = {"schema_version": SCHEMA_VERSION, "status": "native_text",
           "extraction": {"method": "python-docx",
                          "tool": "python-docx %s" % getattr(docx, "__version__", "?")},
           "meta": {},
           "pages": [{"n": 1, "width": None, "height": None, "blocks": blocks,
                      "chars": sum(len(b["text"]) for b in blocks), "status": "native_text",
                      "method": "native", "warnings": pwarn}],
           "assets": assets, "warnings": warnings,
           "source": {"path": name, "name": name, "format": "docx", "sha256": sha,
                      "size": len(raw)}}
    if ok:
        out["questions"] = qs
        out["intro"] = "\n".join(intro).strip()
    return out
